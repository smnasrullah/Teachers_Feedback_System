"""Exports (P17): course, department, university and pending reports as CSV, Excel and PDF.

Django test-client tests. PDF tests skip when WeasyPrint's native libraries are missing.
"""

import csv
import io

import pytest
from django.urls import reverse
from django.utils import timezone, translation
from openpyxl import load_workbook

from apps.analytics import services as analytics
from apps.auditlog.models import AuditEvent
from apps.reports.services import exports
from apps.reports.services.exports import builders, writers
from apps.reports.services.exports.model import Export, Format, Table
from apps.reports.services.exports.pdf import PdfUnavailable, weasyprint_available
from tests.factories import (
    DepartmentFactory,
    claim,
    make_admin,
    make_hod,
    make_student,
    make_teacher,
    sign_in,
)
from tests.test_hod_panel import CLOSED, FOREIGN_ID, OPEN, OWN_ID, build

pytestmark = pytest.mark.django_db

needs_pdf = pytest.mark.skipif(
    not weasyprint_available(), reason="WeasyPrint's native libraries are not installed"
)
FORMULAS = ("=1+1", "+SUM(A1)", "-2+3", "@cmd", "\t=1", "\r=1")


@pytest.fixture
def closed():
    return build(CLOSED)


@pytest.fixture
def open_round():
    return build(OPEN)


@pytest.fixture(autouse=True)
def _english():
    with translation.override("en"):
        yield


def table_rows(response) -> list[list[str]]:
    assert response.content.startswith(b"\xef\xbb\xbf")
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def cells(response) -> set[str]:
    return {cell for row in table_rows(response) for cell in row}


def course_url(world, fmt, offering=None):
    return reverse(
        "reports:export_course", args=[(offering or world.offering).pk, world.window.pk, fmt]
    )


def teacher_of(world):
    return claim(world.offering.teacher)


def exported(action="report.exported"):
    return AuditEvent.objects.filter(action=action)


class TestCourseExport:
    def test_teacher_downloads_csv_with_bom_and_released_figures(self, web, closed):
        closed.submit(closed.offering, count=6)
        sign_in(web, teacher_of(closed))
        response = web.get(course_url(closed, "csv"))
        assert response.status_code == 200
        assert response["Content-Type"].startswith("text/csv")
        assert "attachment" in response["Content-Disposition"]
        assert "CSE201" in cells(response)
        assert "The lecturer shouts at the back row" in cells(response)
        assert "4.0" in cells(response)

    def test_excel_has_one_sheet_per_table(self, web, closed):
        closed.submit(closed.offering, count=6)
        sign_in(web, teacher_of(closed))
        response = web.get(course_url(closed, "xlsx"))
        workbook = load_workbook(io.BytesIO(response.content))
        assert {"Summary", "Questions", "Comments"} <= set(workbook.sheetnames)
        assert "The lecturer shouts at the back row" in {
            cell.value for row in workbook["Comments"].iter_rows() for cell in row
        }

    def test_own_head_and_admin_may_export_others_may_not(self, web, closed):
        closed.submit(closed.offering, count=6)
        url = course_url(closed, "csv")
        for allowed in (closed.hod, make_admin()):
            sign_in(web, allowed)
            assert web.get(url).status_code == 200
        for refused in (
            claim(make_teacher(closed.department)),
            make_hod(DepartmentFactory()),
            claim(make_student(closed.department)),
        ):
            sign_in(web, refused)
            assert web.get(url).status_code == 403
        web.logout()
        assert web.get(url).status_code == 302

    def test_open_round_exports_nothing_but_the_notice(self, web, open_round):
        open_round.submit(open_round.offering, count=6)
        sign_in(web, teacher_of(open_round))
        for fmt in ("csv", "xlsx"):
            assert web.get(course_url(open_round, fmt)).status_code == 200
        response = web.get(course_url(open_round, "csv"))
        assert "Results are released after the round closes." in cells(response)
        assert "The lecturer shouts at the back row" not in cells(response)
        assert "4.0" not in cells(response)

    def test_a_group_below_the_threshold_exports_no_figures_and_no_comments(self, web, closed):
        closed.submit(closed.offering, count=4)
        sign_in(web, teacher_of(closed))
        response = web.get(course_url(closed, "csv"))
        assert "Not enough responses to show this." in cells(response)
        assert "The lecturer shouts at the back row" not in cells(response)
        assert "4.0" not in cells(response)

    def test_comments_come_shuffled_and_without_ids_or_times(self, web, closed):
        for index in range(8):
            closed.submit(closed.offering, count=1, text=f"Comment number {index}")
        sign_in(web, teacher_of(closed))
        response = web.get(course_url(closed, "csv"))
        listed = [
            row[2]
            for row in table_rows(response)
            if len(row) == 3 and row[2].startswith("Comment number")
        ]
        stored = analytics.comment_list(
            closed.window, closed.offering, closed.comment, viewer=closed.hod
        )
        assert listed == [comment.text for comment in stored]
        assert sorted(listed) == [f"Comment number {index}" for index in range(8)]
        assert listed != sorted(listed)
        assert all(
            not comment.token or comment.token not in response.content.decode()
            for comment in stored
        )
        headers = {cell.lower() for row in table_rows(response) for cell in row}
        assert not headers & {"time", "date", "timestamp", "id", "token"}

    def test_every_download_is_audited_without_content(self, web, closed):
        closed.submit(closed.offering, count=6)
        sign_in(web, teacher_of(closed))
        for fmt in ("csv", "xlsx"):
            web.get(course_url(closed, fmt))
        events = list(exported())
        assert [event.metadata["format"] for event in events] == ["csv", "xlsx"]
        assert events[0].metadata["report"] == "course"
        assert events[0].object_id == str(closed.offering.pk)
        assert events[0].department_id == closed.department.pk
        assert "shouts" not in str(events[0].metadata)

    def test_unknown_format_and_offering_are_404(self, web, closed):
        sign_in(web, teacher_of(closed))
        assert web.get(course_url(closed, "csv").replace("csv", "docx")).status_code == 404
        missing = reverse("reports:export_course", args=[9999, closed.window.pk, "csv"])
        assert web.get(missing).status_code == 404


class TestDepartmentExport:
    def url(self, fmt="csv", **params):
        query = "&".join(f"{key}={value}" for key, value in params.items())
        return reverse("reports:export_department", args=[fmt]) + (f"?{query}" if query else "")

    def test_head_exports_the_ranking_and_the_threshold_holds(self, web, closed):
        closed.submit(closed.offering, count=6)
        closed.submit(closed.colleague_offering, count=4, rating=1)
        sign_in(web, closed.hod)
        response = web.get(self.url(below="4.5"))
        rows = table_rows(response)
        by_code = {row[1]: row for row in rows if len(row) == 12 and row[1].startswith("CSE")}
        assert by_code["CSE201"][0] == "1"
        assert by_code["CSE201"][5] == "4.0"
        assert by_code["CSE201"][9] == "Yes"
        assert by_code["CSE202"][5] == ""
        assert by_code["CSE202"][11] == "Not enough responses to show this."
        assert "EEE301" not in cells(response)

    def test_open_round_has_no_figures(self, web, open_round):
        open_round.submit(open_round.offering, count=6)
        sign_in(web, open_round.hod)
        rows = table_rows(web.get(self.url()))
        course = next(row for row in rows if len(row) == 12 and row[1] == "CSE201")
        assert course[0] == "" and course[5] == ""

    def test_excel_and_audit(self, web, closed):
        closed.submit(closed.offering, count=6)
        sign_in(web, closed.hod)
        response = web.get(self.url("xlsx"))
        assert "Courses" in load_workbook(io.BytesIO(response.content)).sheetnames
        event = exported().get()
        assert event.metadata["report"] == "department"
        assert event.department_id == closed.department.pk

    def test_admin_must_name_a_department_and_teachers_are_refused(self, web, closed):
        sign_in(web, make_admin())
        assert web.get(self.url()).status_code == 400
        assert web.get(self.url(department=closed.department.pk)).status_code == 200
        assert web.get(self.url(department=9999)).status_code == 400
        sign_in(web, claim(make_teacher(closed.department)))
        assert web.get(self.url()).status_code == 403

    def test_head_cannot_reach_another_department_by_parameter(self, web, closed):
        closed.submit(closed.foreign_offering, count=6)
        sign_in(web, closed.hod)
        response = web.get(self.url(department=closed.foreign_offering.course.department_id))
        assert "EEE301" not in cells(response)


class TestUniversityExport:
    def url(self, fmt="csv", **params):
        query = "&".join(f"{key}={value}" for key, value in params.items())
        return reverse("reports:export_university", args=[fmt]) + (f"?{query}" if query else "")

    def test_only_the_admin_exports_it(self, web, closed):
        closed.submit(closed.offering, count=6)
        for refused in (
            closed.hod,
            claim(make_teacher(closed.department)),
            claim(make_student(closed.department)),
        ):
            sign_in(web, refused)
            assert web.get(self.url()).status_code == 403
        sign_in(web, make_admin())
        response = web.get(self.url())
        assert response.status_code == 200
        assert "CSE201" in cells(response)

    def test_filters_and_hidden_groups(self, web, closed):
        closed.submit(closed.offering, count=6)
        closed.submit(closed.foreign_offering, count=3)
        sign_in(web, make_admin())
        everything = cells(web.get(self.url()))
        assert "EEE301" not in everything
        narrowed = web.get(self.url(department=closed.department.pk))
        assert "CSE201" in cells(narrowed)
        assert web.get(self.url(department="abc")).status_code == 400

    def test_open_round_is_flagged_and_unreleased(self, web, open_round):
        open_round.submit(open_round.offering, count=6)
        sign_in(web, make_admin())
        response = web.get(self.url())
        assert "Results are released after the round closes." in cells(response)
        assert "CSE201" not in cells(response)

    def test_excel_and_audit(self, web, closed):
        closed.submit(closed.offering, count=6)
        sign_in(web, make_admin())
        response = web.get(self.url("xlsx"))
        assert "Departments" in load_workbook(io.BytesIO(response.content)).sheetnames
        assert exported().get().metadata["report"] == "university"


class TestPendingExport:
    def url(self, fmt="csv", **params):
        query = "&".join(f"{key}={value}" for key, value in params.items())
        return reverse("reports:export_pending", args=[fmt]) + (f"?{query}" if query else "")

    def test_head_gets_own_department_ids_only(self, web, closed):
        closed.enroll(closed.department, external_id=OWN_ID)
        closed.enroll(closed.foreign_offering.course.department, external_id=FOREIGN_ID)
        closed.enroll(
            closed.foreign_offering.course.department,
            offering=closed.offering,
            external_id="2209888",
        )
        sign_in(web, closed.hod)
        response = web.get(self.url())
        assert OWN_ID in cells(response)
        assert FOREIGN_ID not in cells(response)
        assert "2209888" not in cells(response)
        counted = next(
            row
            for row in table_rows(response)
            if row and row[0] == "Students of other departments pending"
        )
        assert counted[1] == "2"
        assert exported().get().metadata["report"] == "pending"

    def test_submitted_students_are_not_listed(self, web, closed):
        closed.submit(closed.offering, count=1)
        closed.enroll(closed.department, external_id=OWN_ID)
        sign_in(web, closed.hod)
        assert len([c for c in cells(web.get(self.url())) if c.startswith("2101")]) == 1

    def test_course_filter_and_foreign_course_refused(self, web, closed):
        closed.enroll(closed.department, external_id=OWN_ID)
        closed.enroll(closed.department, offering=closed.colleague_offering, external_id="2101999")
        sign_in(web, closed.hod)
        only = cells(web.get(self.url(course=closed.offering.pk)))
        assert OWN_ID in only and "2101999" not in only
        assert web.get(self.url(course=closed.foreign_offering.pk)).status_code == 403

    def test_admin_names_a_department_and_sees_its_ids(self, web, closed):
        closed.enroll(closed.department, external_id=OWN_ID)
        sign_in(web, make_admin())
        assert web.get(self.url()).status_code == 400
        assert OWN_ID in cells(web.get(self.url(department=closed.department.pk)))

    def test_teachers_students_and_other_heads_are_refused(self, web, closed):
        closed.enroll(closed.department, external_id=OWN_ID)
        other_head = make_hod(DepartmentFactory())
        for refused in (
            claim(make_teacher(closed.department)),
            claim(make_student(closed.department)),
        ):
            sign_in(web, refused)
            assert web.get(self.url()).status_code == 403
        sign_in(web, other_head)
        assert web.get(self.url()).status_code == 404

    def test_excel_has_no_time_columns(self, web, closed):
        closed.enroll(closed.department, external_id=OWN_ID)
        sign_in(web, closed.hod)
        sheet = load_workbook(io.BytesIO(web.get(self.url("xlsx")).content))["Pending students"]
        header = [cell.value for cell in sheet[sheet.max_row - 1 if sheet.max_row > 2 else 1]]
        assert not {"Time", "Date"} & set(header)


class TestFormulaInjection:
    def export(self):
        rows = tuple((value, f"x{value}", 5, None) for value in FORMULAS)
        return Export(
            "k",
            "n",
            "=title",
            "+subtitle",
            (Table("=sheet", ("-h", "a", "b", "c"), rows),),
            notices=("@note",),
        )

    def test_csv_cells_are_neutralised(self):
        text = writers.to_csv(self.export()).decode("utf-8-sig")
        parsed = [cell for row in csv.reader(io.StringIO(text)) for cell in row]
        for value in ("=title", "+subtitle", "@note", "=sheet", "-h", *FORMULAS):
            assert value not in parsed
            assert "'" + value in parsed
        assert "5" in parsed

    def test_excel_cells_are_text_not_formulas(self):
        workbook = load_workbook(io.BytesIO(writers.to_xlsx(self.export())))
        assert workbook.sheetnames == ["=sheet"]
        for row in workbook.active.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                if isinstance(cell.value, str):
                    assert cell.value[0] not in "=+-@\t\r"

    def test_a_comment_that_looks_like_a_formula_is_defused_end_to_end(self, web, closed):
        closed.submit(closed.offering, count=6, text='=HYPERLINK("http://evil","click")')
        sign_in(web, teacher_of(closed))
        assert '\'=HYPERLINK("http://evil","click")' in cells(web.get(course_url(closed, "csv")))
        sheet = load_workbook(io.BytesIO(web.get(course_url(closed, "xlsx")).content))["Comments"]
        values = [
            cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)
        ]
        assert '\'=HYPERLINK("http://evil","click")' in values

    def test_control_characters_do_not_break_excel(self):
        export = Export("k", "n", "t", "s", (Table("T", ("a",), (("bad\x00\x07text",),)),))
        sheet = load_workbook(io.BytesIO(writers.to_xlsx(export))).active
        assert "badtext" in {cell.value for row in sheet.iter_rows() for cell in row}


class TestPdf:
    @needs_pdf
    def test_course_pdf_embeds_the_bangla_font(self, web, closed):
        closed.submit(closed.offering, count=6, text="চমৎকার পাঠদান")
        closed.offering.course.title_bn = "প্রোগ্রামিং"
        closed.offering.course.save()
        sign_in(web, teacher_of(closed))
        response = web.get(course_url(closed, "pdf"), HTTP_ACCEPT_LANGUAGE="bn")
        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")
        assert b"HindSiliguri" in response.content
        assert exported().get().metadata["format"] == "pdf"

    @needs_pdf
    def test_every_report_makes_a_pdf(self, web, closed):
        closed.submit(closed.offering, count=6)
        closed.enroll(closed.department, external_id=OWN_ID)
        sign_in(web, closed.hod)
        for url in (
            course_url(closed, "pdf"),
            reverse("reports:export_department", args=["pdf"]),
            reverse("reports:export_pending", args=["pdf"]),
        ):
            assert web.get(url).content.startswith(b"%PDF")
        sign_in(web, make_admin())
        assert web.get(reverse("reports:export_university", args=["pdf"])).content.startswith(
            b"%PDF"
        )

    def test_missing_weasyprint_answers_503_and_leaves_no_audit_event(
        self, web, closed, monkeypatch
    ):
        def broken(_export):
            raise PdfUnavailable

        monkeypatch.setitem(exports.WRITERS, Format.PDF, broken)
        closed.submit(closed.offering, count=6)
        sign_in(web, teacher_of(closed))
        assert web.get(course_url(closed, "pdf")).status_code == 503
        assert not exported().exists()

    def test_pdf_page_reuses_the_svg_bars_and_has_no_script(self, closed):
        from django.template.loader import render_to_string

        closed.submit(closed.offering, count=6)
        export = builders.teacher_course(closed.offering, closed.window, closed.hod, timezone.now())
        html = render_to_string("reports/pdf/export.html", {"export": export})
        assert "<svg" in html and "bar-val" in html
        assert "<script" not in html
        assert "The lecturer shouts at the back row" in html


class TestLinks:
    def test_report_pages_offer_the_downloads_once_results_are_open(self, web, closed):
        closed.submit(closed.offering, count=6)
        sign_in(web, teacher_of(closed))
        assert (
            course_url(closed, "csv")
            in web.get(closed.report_url(closed.offering)).content.decode()
        )
        sign_in(web, closed.hod)
        assert (
            reverse("reports:export_department", args=["xlsx"])
            in web.get(reverse("accounts:dashboard_hod")).content.decode()
        )
        assert (
            reverse("reports:export_pending", args=["pdf"])
            in web.get(reverse("reports:hod_pending")).content.decode()
        )
        sign_in(web, make_admin())
        assert (
            reverse("reports:export_university", args=["csv"])
            in web.get(reverse("reports:analytics")).content.decode()
        )


def test_file_names_are_safe_slugs():
    export = Export("course", "CSE 201/../Round 1", "t", "s", ())
    assert exports.filename(export, Format.CSV) == "course-cse-201round-1.csv"
