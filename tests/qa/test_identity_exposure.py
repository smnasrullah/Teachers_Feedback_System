"""BLUEPRINT §5.3 row 1 across the whole product: a Student ID or full name reaches only the Admin
and the head of that student's own department, on pages, error pages and exports alike."""

import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.core.services.demo import constants as demo
from tests.qa.conftest import all_routes, client_for, sample_kwargs

CSE_OTHERS = tuple(zip(demo.CSE_STUDENT_IDS[1:], demo.CSE_STUDENT_NAMES[1:], strict=True))
EEE = tuple(zip(demo.EEE_STUDENT_IDS, demo.EEE_STUDENT_NAMES, strict=True))
# What each role must never see: (Student ID, full name) pairs.
FORBIDDEN = {
    demo.STUDENT: (*CSE_OTHERS, *EEE, (demo.CSE_STUDENT_IDS[0], None)),
    demo.TEACHER: (*CSE_OTHERS, *EEE, (demo.CSE_STUDENT_IDS[0], demo.CSE_STUDENT_NAMES[0])),
    demo.HOD: EEE,
}


def _text(response) -> str:
    if response.streaming:
        return b"".join(response.streaming_content).decode(errors="replace")
    content_type = response.get("Content-Type", "")
    if "spreadsheetml" in content_type:
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        return "\n".join(
            str(cell)
            for sheet in book.worksheets
            for row in sheet.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        )
    return response.content.decode(errors="replace")


@pytest.mark.parametrize("role", list(FORBIDDEN))
def test_no_page_or_export_shows_an_id_or_name_the_role_may_not_see(seeded, role):
    client = client_for(role)
    user = User.objects.get(username=role)
    leaks = []
    for route in all_routes():
        if "<path:object_id>" in route.pattern or route.name.startswith("admin:"):
            continue
        for kwargs in sample_kwargs(route.name, user):
            url = reverse(route.name, kwargs=kwargs)
            body = _text(client.get(url))
            for external_id, name in FORBIDDEN[role]:
                leaks += [
                    f"{url}: {value}" for value in (external_id, name) if value and value in body
                ]
    assert not leaks, "\n".join(leaks)


def test_the_detector_sees_ids_where_they_are_allowed(seeded):
    """Control: the same reading finds IDs on pages that may show them."""
    admin = _text(
        client_for(demo.ADMIN).get(
            reverse("accounts:registry_list"), {"q": demo.EEE_STUDENT_IDS[0]}
        )
    )
    head = _text(client_for(demo.HOD).get(reverse("reports:export_pending", args=["xlsx"])))

    assert demo.EEE_STUDENT_IDS[0] in admin
    assert any(external_id in head for external_id in demo.CSE_STUDENT_IDS)
