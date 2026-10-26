"""CSV (UTF-8 with BOM, so Excel shows Bangla) and Excel files.

Every text cell goes through ``safe_cell``: one that starts with ``= + - @`` would run as a
formula when opened, so it is prefixed with an apostrophe. Numbers are written as numbers.
"""

import csv
import io
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

from apps.core.csv_safe import safe_cell
from apps.reports.services.exports.model import Export

BOM = "﻿"
SHEET_TITLE_MAX = 31
SHEET_TITLE_FORBIDDEN = re.compile(r"[\/*?:\[\]]")
COLUMN_MIN = 10
COLUMN_MAX = 60
WRAP_FROM = 40


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return safe_cell(ILLEGAL_CHARACTERS_RE.sub("", value))
    return value


def to_csv(export: Export) -> bytes:
    """One file: the title, then each table under its own heading, separated by a blank line."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([_cell(export.title)])
    writer.writerow([_cell(export.subtitle)])
    for notice in export.notices:
        writer.writerow([_cell(notice)])
    for table in export.tables:
        writer.writerow([])
        writer.writerow([_cell(table.title)])
        writer.writerow([_cell(header) for header in table.headers])
        for row in table.rows:
            writer.writerow([_cell(value) for value in row])
    return (BOM + buffer.getvalue()).encode("utf-8")


def _sheet_title(title: str, used: set[str]) -> str:
    base = SHEET_TITLE_FORBIDDEN.sub(" ", title).strip()[:SHEET_TITLE_MAX] or "Sheet"
    candidate, number = base, 1
    while candidate.lower() in used:
        number += 1
        suffix = f" {number}"
        candidate = base[: SHEET_TITLE_MAX - len(suffix)] + suffix
    used.add(candidate.lower())
    return candidate


def _put(sheet, row: int, values) -> None:
    for column, value in enumerate(values, start=1):
        sheet.cell(row=row, column=column, value=_cell(value))


def to_xlsx(export: Export) -> bytes:
    """One sheet per table; the first sheet also carries the title and the notices."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for index, table in enumerate(export.tables):
        sheet = workbook.create_sheet(_sheet_title(table.title, used))
        row = 0
        if index == 0:
            for line in (export.title, export.subtitle, *export.notices):
                row += 1
                _put(sheet, row, [line])
            sheet["A1"].font = Font(bold=True, size=14)
            row += 1
        row += 1
        head_row = row
        _put(sheet, row, table.headers)
        for cell in sheet[head_row]:
            cell.font = Font(bold=True)
        for values in table.rows:
            row += 1
            _put(sheet, row, values)
        for column in sheet.iter_cols(min_row=head_row, max_row=row):
            longest = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(
                max(longest + 2, COLUMN_MIN), COLUMN_MAX
            )
            if longest >= WRAP_FROM:
                for cell in column[1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
